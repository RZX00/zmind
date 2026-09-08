from __future__ import annotations

import asyncio
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Optional, Tuple

import chardet
import PyPDF2
import docx
from PIL import Image

from core.services.supabase import DBConnection
from core.services.deepseek_ocr import deepseek_ocr_pdf_bytes
from core.services.r2_storage import get_object_bytes
from core.services.transcription import transcribe_audio_bytes
from core.utils.config import config
from core.utils.logger import logger


TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".json", ".yaml", ".yml", ".xml", ".csv", ".tsv",
    ".log", ".ini", ".cfg", ".conf",
}
DOCX_EXTS = {".docx"}
PDF_EXTS = {".pdf"}
PPTX_EXTS = {".pptx"}
XLSX_EXTS = {".xlsx", ".xls"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}
VIDEO_EXTS = {".mp4", ".mov", ".webm", ".m4v", ".mkv", ".avi"}
AUDIO_EXTS = {".mp3", ".wav", ".ogg", ".flac", ".m4a", ".mp4", ".webm"}


def _normalize_text(text: str, max_chars: int) -> Tuple[str, bool]:
    if max_chars and len(text) > max_chars:
        return text[:max_chars], True
    return text, False


def _decode_text_bytes(data: bytes) -> str:
    detected = chardet.detect(data[:4096])
    encoding = detected.get("encoding") or "utf-8"
    try:
        return data.decode(encoding)
    except Exception:
        return data.decode("utf-8", errors="replace")


def _extract_text_from_pdf(data: bytes) -> str:
    reader = PyPDF2.PdfReader(BytesIO(data))
    chunks = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if txt:
            chunks.append(txt)
    return "\n\n".join(chunks)


def _extract_text_from_docx(data: bytes) -> str:
    doc = docx.Document(BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _extract_text_from_pptx(data: bytes, max_chars: int) -> str:
    try:
        from pptx import Presentation
    except Exception:
        return ""
    prs = Presentation(BytesIO(data))
    parts = []
    total = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "") or ""
            if not text:
                continue
            parts.append(text)
            total += len(text)
            if max_chars and total >= max_chars:
                break
        if max_chars and total >= max_chars:
            break
    return "\n".join(parts)


def _extract_text_from_xlsx(data: bytes, max_chars: int) -> str:
    try:
        import openpyxl
    except Exception:
        return ""
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    parts = []
    total = 0
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join("" if v is None else str(v) for v in row)
            if row_text.strip():
                parts.append(row_text)
                total += len(row_text)
                if max_chars and total >= max_chars:
                    break
        if max_chars and total >= max_chars:
            break
    return "\n".join(parts)


def _is_textual(content_type: str, ext: str) -> bool:
    ct = (content_type or "").lower()
    if ct.startswith("text/"):
        return True
    if ct in (
        "application/json",
        "application/xml",
        "text/csv",
        "application/csv",
        "application/x-yaml",
        "application/yaml",
    ):
        return True
    return ext in TEXT_EXTS


def _is_image(content_type: str, ext: str) -> bool:
    return (content_type or "").lower().startswith("image/") or ext in IMAGE_EXTS


def _is_audio_or_video(content_type: str, ext: str) -> Tuple[bool, bool]:
    ct = (content_type or "").lower()
    if ct.startswith("audio/"):
        return True, False
    if ct.startswith("video/"):
        return False, True
    if ext in AUDIO_EXTS:
        return True, False
    if ext in VIDEO_EXTS:
        return False, True
    return False, False


def _should_run_deepseek_ocr(
    metadata: dict,
    extracted_text: Optional[str],
    content_type: str,
    ext: str,
) -> bool:
    if not config.DEEPSEEK_OCR_ENABLED:
        return False
    is_pdf = ext == ".pdf" or (content_type or "").lower() == "application/pdf"
    if not is_pdf:
        return False
    if metadata.get("ocr_requested") or metadata.get("force_ocr"):
        return True
    if config.DEEPSEEK_OCR_AUTO_ON_EMPTY_PDF and not (extracted_text or "").strip():
        return True
    return False


async def _download_bytes(record: dict, client) -> Optional[bytes]:
    provider = record.get("storage_provider") or "supabase"
    key = record.get("storage_path")
    if not key:
        return None
    if provider == "r2":
        return await asyncio.to_thread(get_object_bytes, key)
    return None


async def process_file_upload(
    file_id: str,
    account_id: Optional[str] = None,
    file_bytes: Optional[bytes] = None,
) -> None:
    client = await DBConnection().client
    query = client.table("file_uploads") \
        .select("id, account_id, user_id, storage_provider, storage_bucket, storage_path, content_type, original_filename, file_size, status, metadata, source") \
        .eq("id", file_id)
    if account_id:
        query = query.eq("account_id", account_id)
    record_resp = await query.limit(1).execute()
    if not record_resp.data:
        return
    record = record_resp.data[0]

    status = (record.get("status") or "").lower()
    if status in ("failed", "deleted"):
        return

    metadata = record.get("metadata") or {}
    if isinstance(metadata, str):
        metadata = {}
    force_ocr = bool(metadata.get("ocr_requested") or metadata.get("force_ocr"))
    if metadata.get("processing_status") == "ready" and not force_ocr:
        return

    file_size = int(record.get("file_size") or 0)
    if file_size <= 0:
        return

    max_bytes = int(config.FILE_PROCESSING_MAX_BYTES or 0)
    if max_bytes and file_size > max_bytes:
        metadata.update({
            "processing_status": "skipped",
            "processing_reason": "file_too_large",
            "processed_at": datetime.now(timezone.utc).isoformat(),
        })
        await client.table("file_uploads").update({
            "metadata": metadata,
        }).eq("id", file_id).execute()
        return

    try:
        metadata.update({
            "processing_status": "processing",
            "processing_started_at": datetime.now(timezone.utc).isoformat(),
        })
        await client.table("file_uploads").update({
            "metadata": metadata,
            "status": "processing",
        }).eq("id", file_id).execute()
    except Exception:
        pass

    filename = record.get("original_filename") or "file"
    content_type = record.get("content_type") or "application/octet-stream"
    ext = Path(filename).suffix.lower()

    try:
        data = file_bytes if file_bytes is not None else await _download_bytes(record, client)
        if not data:
            raise RuntimeError("missing file bytes")

        extracted_text: Optional[str] = None
        extracted_type: Optional[str] = None
        image_meta: Optional[dict] = None
        source = (record.get("source") or "").lower()
        skip_audio_transcription = source == "canvas-audio"

        if _is_textual(content_type, ext):
            extracted_text = _decode_text_bytes(data)
            extracted_type = "text"
        elif ext in PDF_EXTS or content_type == "application/pdf":
            extracted_text = _extract_text_from_pdf(data)
            extracted_type = "text"
            if _should_run_deepseek_ocr(metadata, extracted_text, content_type, ext):
                try:
                    ocr_text, ocr_meta = await deepseek_ocr_pdf_bytes(
                        data,
                        max_pages=int(config.DEEPSEEK_OCR_MAX_PAGES or 6),
                        dpi=int(config.DEEPSEEK_OCR_DPI or 200),
                        max_image_dim=int(config.DEEPSEEK_OCR_MAX_IMAGE_DIM or 1600),
                    )
                    ocr_meta.update({
                        "provider": "deepseek-ocr",
                        "requested": bool(metadata.get("ocr_requested") or metadata.get("force_ocr")),
                        "processed_at": datetime.now(timezone.utc).isoformat(),
                    })
                    metadata["ocr"] = ocr_meta
                    if ocr_text:
                        extracted_text = ocr_text
                        extracted_type = "ocr"
                except Exception as exc:
                    metadata["ocr"] = {
                        "status": "failed",
                        "provider": "deepseek-ocr",
                        "error": str(exc),
                        "processed_at": datetime.now(timezone.utc).isoformat(),
                    }
        elif ext in DOCX_EXTS or content_type in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",):
            extracted_text = _extract_text_from_docx(data)
            extracted_type = "text"
        elif ext in PPTX_EXTS or content_type in ("application/vnd.openxmlformats-officedocument.presentationml.presentation",):
            extracted_text = _extract_text_from_pptx(data, config.FILE_PROCESSING_TEXT_MAX_CHARS)
            extracted_type = "text"
        elif ext in XLSX_EXTS or content_type in ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",):
            extracted_text = _extract_text_from_xlsx(data, config.FILE_PROCESSING_TEXT_MAX_CHARS)
            extracted_type = "text"
        else:
            is_audio, is_video = _is_audio_or_video(content_type, ext)
            if is_audio or is_video:
                if skip_audio_transcription:
                    extracted_text = None
                    extracted_type = None
                else:
                    transcript = await transcribe_audio_bytes(
                        data,
                        filename=filename,
                        content_type=content_type,
                        user_id=str(record.get("user_id") or ""),
                    )
                    extracted_text = transcript
                    extracted_type = "transcript"

        if _is_image(content_type, ext):
            try:
                with Image.open(BytesIO(data)) as img:
                    image_meta = {
                        "width": img.width,
                        "height": img.height,
                        "format": img.format,
                    }
            except Exception:
                image_meta = None

        max_chars = int(config.FILE_PROCESSING_TEXT_MAX_CHARS or 0)
        truncated = False
        if extracted_text:
            extracted_text, truncated = _normalize_text(extracted_text, max_chars)

        if metadata.get("ocr_requested"):
            metadata["ocr_requested"] = False
        if metadata.get("force_ocr"):
            metadata["force_ocr"] = False

        metadata.update({
            "processing_status": "ready" if extracted_text or image_meta else "skipped",
            "processed_at": datetime.now(timezone.utc).isoformat(),
        })
        if extracted_text:
            metadata.update({
                "extracted_text": extracted_text,
                "extracted_text_truncated": truncated,
                "extracted_text_type": extracted_type or "text",
            })
        if image_meta:
            metadata["image_meta"] = image_meta

        await client.table("file_uploads").update({
            "metadata": metadata,
            "status": "ready" if extracted_text or image_meta else record.get("status"),
        }).eq("id", file_id).execute()
    except Exception as e:
        metadata.update({
            "processing_status": "failed",
            "processing_error": str(e),
            "processed_at": datetime.now(timezone.utc).isoformat(),
        })
        try:
            await client.table("file_uploads").update({
                "metadata": metadata,
                "status": "uploaded",
            }).eq("id", file_id).execute()
        except Exception:
            pass
        logger.warning(f"[files] processing failed for {file_id}: {e}")
